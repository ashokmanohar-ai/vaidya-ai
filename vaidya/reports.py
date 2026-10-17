"""Deterministic DOCX renderer for clinician-reviewed drafts."""
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from pathlib import Path
import io
import re


def report_text(findings, instructions, flags):
    lines=['HISTOPATHOLOGY — DRAFT / NOT VALIDATED','DICTATED OBSERVATIONS']
    lines += [f'• {x["text"]}' for x in findings]
    if flags:
        lines += ['','REQUIRES DOCTOR REVIEW']+[f'• {f}' for f in flags]
    return '\n'.join(lines)


def render_docx(job,patient,hospital,findings,instructions,flags,images,output,approved_skills=""):
    doc=Document()
    section=doc.sections[0]; section.header_distance=Inches(.35)
    header=section.header.paragraphs[0]
    header.text=hospital['logo_text']
    header.style='Title'
    sub=doc.add_paragraph(hospital['address']);sub.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading('HISTOPATHOLOGY REPORT',level=1)
    warn=doc.add_paragraph('AI-ASSISTED DRAFT · NOT FOR CLINICAL USE · DOCTOR APPROVAL REQUIRED')
    for r in warn.runs: r.bold=True;r.font.color.rgb=RGBColor(190,46,30)
    t=doc.add_table(rows=0,cols=2);t.style='Light Shading Accent 1';t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for a,b in [('Patient reference',patient['code']),('Patient name',patient['name']),('Age / Sex',f"{patient['age']} / {patient['sex']}"),
                ('Report type','Histopathology'),('Job reference',job['id'][:12]),('Status','DRAFT – DOCTOR REVIEW')]:
        row=t.add_row().cells;row[0].text=a;row[1].text=str(b)
    doc.add_heading('Dictated observations',level=2)
    bold_targets=[]
    for ins in instructions:
        match=re.search(r'\b(?:please )?(?:bold|highlight)\s+(.+?)(?:\s+in\s+the\s+report)?[.!]?$' , ins['instruction'], flags=re.I)
        if match: bold_targets.append(match.group(1).strip(' .!'))
    for fact in findings:
        sentence=str(fact['text'])
        paragraph=doc.add_paragraph(style='List Bullet')
        highlighted=False
        for target in bold_targets:
            m=re.search(re.escape(target),sentence,flags=re.I)
            if m:
                paragraph.add_run(sentence[:m.start()]);paragraph.add_run(sentence[m.start():m.end()]).bold=True
                paragraph.add_run(sentence[m.end():]);highlighted=True;break
        if not highlighted: paragraph.add_run(sentence)
    if not findings:doc.add_paragraph('[No dictated observation extracted – doctor review required]')
    if images:
        doc.add_heading('Uploaded reference images',level=2)
        for idx,p in enumerate(images,1):
            doc.add_paragraph(f'Image {idx} · uploaded slide reference; no AI diagnosis performed')
            try:doc.add_picture(str(p),width=Inches(3.7))
            except Exception:doc.add_paragraph('[Image could not be embedded]')
    if flags:
        doc.add_heading('Items for clinician clarification',level=2)
        for flag in flags:doc.add_paragraph(flag,style='List Bullet')
    if 'APPROVED:' in approved_skills and 'signature' in approved_skills.lower():
        doc.add_paragraph('\nClinician signature: _______________________     Date: __________')
    footer=section.footer.paragraphs[0];footer.text='VAIDYA AI · Synthetic pilot demonstration · Not a validated medical document'
    Path(output).parent.mkdir(parents=True,exist_ok=True)
    doc.save(str(output))
    return str(output)


def final_docx(text,job,patient,hospital,output):
    doc=Document()
    doc.add_heading(hospital['logo_text'],0)
    doc.add_paragraph('CLINICIAN-APPROVED SYNTHETIC DEMONSTRATION')
    doc.add_heading('Histopathology Report',1)
    doc.add_paragraph(f"Patient: {patient['name']} ({patient['code']})")
    for l in text.splitlines():
        if l.strip() in ("HISTOPATHOLOGY — DRAFT / NOT VALIDATED", "DICTATED OBSERVATIONS"): continue
        doc.add_paragraph(l)
    doc.add_paragraph(f"Approved by demo clinician · job {job['id'][:12]}")
    doc.save(str(output))
